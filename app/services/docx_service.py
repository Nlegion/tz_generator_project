from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os
from datetime import datetime
import structlog

from app.core.config import settings
from app.utils.formatting import parse_markdown_to_docx_structure, DocxBlock, BlockType

logger = structlog.get_logger(__name__)


class DocxGenerator:
    def __init__(self, template_path: str = None):
        self.template_path = template_path or settings.TEMPLATE_PATH

    def create_document(self, tz_data: dict, output_path: str) -> str:
        """
        Создание .docx документа из текста ТЗ

        Args:
            tz_data: Данные ТЗ (текст и метаданные)
            output_path: Путь для сохранения файла

        Returns:
            Путь к сохраненному файлу
        """
        try:
            # Создаем или загружаем документ из шаблона
            if self.template_path and os.path.exists(self.template_path):
                doc = Document(self.template_path)
                logger.debug('template_loaded', template_path=self.template_path)
            else:
                doc = Document()
                self._setup_default_styles(doc)
                logger.warning('template_not_found', template_path=self.template_path, using_default=True)

            # Добавляем заголовок
            self._add_title(doc, 'Техническое задание')

            # Добавляем метаданные
            self._add_metadata(doc, tz_data.get('metadata', {}))

            # Парсим и добавляем текст через улучшенный парсер
            text = tz_data.get('text', '')
            blocks = parse_markdown_to_docx_structure(text)
            self._add_blocks(doc, blocks)

            # Сохраняем документ
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            doc.save(output_path)

            logger.info('document_created', output_path=output_path, blocks_count=len(blocks))
            return output_path

        except (OSError, PermissionError) as e:
            msg = f'Ошибка создания документа: {e}'
            logger.error('document_creation_error', error=str(e), output_path=output_path)
            raise OSError(msg) from e

    def _setup_default_styles(self, doc):
        """Настройка стилей по умолчанию"""
        styles = doc.styles

        # Стиль заголовка 1
        heading1 = styles.add_style('TZHeading1', WD_STYLE_TYPE.PARAGRAPH)
        heading1.font.name = 'Times New Roman'
        heading1.font.size = Pt(16)
        heading1.font.bold = True
        heading1.paragraph_format.space_after = Pt(12)

        # Стиль заголовка 2
        heading2 = styles.add_style('TZHeading2', WD_STYLE_TYPE.PARAGRAPH)
        heading2.font.name = 'Times New Roman'
        heading2.font.size = Pt(14)
        heading2.font.bold = True
        heading2.paragraph_format.space_before = Pt(6)
        heading2.paragraph_format.space_after = Pt(6)

        # Стиль основного текста
        normal = styles['Normal']
        normal.font.name = 'Times New Roman'
        normal.font.size = Pt(12)
        normal.paragraph_format.line_spacing = 1.5

        # Стиль списка
        list_style = styles.add_style('TZList', WD_STYLE_TYPE.PARAGRAPH)
        list_style.font.name = 'Times New Roman'
        list_style.font.size = Pt(12)
        list_style.paragraph_format.left_indent = Inches(0.25)
        list_style.paragraph_format.first_line_indent = Inches(-0.25)

    def _add_title(self, doc, title):
        """Добавление основного заголовка"""
        try:
            p = doc.add_paragraph(title)
            p.style = 'TZTitle'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except KeyError:
            # Если стиль TZTitle не найден, используем TZHeading1
            p = doc.add_paragraph(title)
            try:
                p.style = 'TZHeading1'
            except KeyError:
                p.style = 'Heading 1'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_metadata(self, doc, metadata):
        """Добавление метаданных"""
        meta_text = f"""
        Дата создания: {datetime.now().strftime('%d.%m.%Y')}
        Тип проекта: {metadata.get('project_type', 'Не указан')}
        Стиль документа: {metadata.get('style', 'Стандартный')}
        """

        p = doc.add_paragraph(meta_text)
        p.style = 'Normal'

    def _add_blocks(self, doc, blocks: list[DocxBlock]):
        """Добавление блоков в документ"""
        for block in blocks:
            if block.block_type == BlockType.EMPTY:
                doc.add_paragraph()
            elif block.block_type == BlockType.HEADING1:
                self._add_heading(doc, block.content, level=1)
            elif block.block_type == BlockType.HEADING2:
                self._add_heading(doc, block.content, level=2)
            elif block.block_type == BlockType.HEADING3:
                self._add_heading(doc, block.content, level=3)
            elif block.block_type == BlockType.LIST_ITEM:
                self._add_list_item(doc, block.content)
            elif block.block_type == BlockType.NUMBERED_ITEM:
                self._add_numbered_item(doc, block.content)
            elif block.block_type == BlockType.PARAGRAPH:
                self._add_paragraph(doc, block.content)

    def _add_heading(self, doc, text, level):
        if level == 1:
            style_name = 'TZHeading1'
        elif level == 2:
            style_name = 'TZHeading2'
        else:
            style_name = 'TZHeading2'  # Fallback для level 3
        
        try:
            p = doc.add_paragraph(text)
            p.style = style_name
        except KeyError:
            # Если стиль не найден, используем стандартный
            p = doc.add_paragraph(text)
            p.style = 'Heading 1' if level == 1 else 'Heading 2'

    def _add_paragraph(self, doc, text):
        doc.add_paragraph(text).style = 'Normal'

    def _add_list_item(self, doc, text):
        p = doc.add_paragraph(text, style='TZList')
        p.paragraph_format.left_indent = Inches(0.25)

    def _add_numbered_item(self, doc, text):
        # Для нумерованных списков используем стандартные стили Word
        p = doc.add_paragraph(text, style='List Number')