import pytest
import os
from pathlib import Path
from docx import Document

from app.utils.file_handlers import load_examples, get_relevant_examples


def test_load_examples_txt(temp_examples_dir, sample_tz_text):
    """Тест загрузки примеров из .txt файлов"""
    # Создаем тестовый файл
    file_path = os.path.join(temp_examples_dir, 'example1.txt')
    with open(file_path, 'w', encoding='utf-8') as f:
        f.write(sample_tz_text)
    
    examples = load_examples(temp_examples_dir)
    
    assert len(examples) == 1
    # Проверяем, что содержимое загружено (может быть с нормализацией пробелов)
    assert len(examples[0]) > 0
    assert 'Техническое задание' in examples[0] or 'Введение' in examples[0]


def test_load_examples_docx(temp_examples_dir):
    """Тест загрузки примеров из .docx файлов"""
    # Создаем тестовый .docx файл
    file_path = os.path.join(temp_examples_dir, 'example1.docx')
    doc = Document()
    doc.add_paragraph('Пример ТЗ из docx')
    doc.add_paragraph('Вторая строка')
    doc.save(file_path)
    
    examples = load_examples(temp_examples_dir)
    
    assert len(examples) == 1
    assert 'Пример ТЗ из docx' in examples[0]
    assert 'Вторая строка' in examples[0]


def test_load_examples_multiple_files(temp_examples_dir, sample_tz_text):
    """Тест загрузки нескольких примеров"""
    # Создаем несколько файлов
    for i in range(3):
        file_path = os.path.join(temp_examples_dir, f'example{i}.txt')
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(f'{sample_tz_text}\n\nFile {i}')
    
    examples = load_examples(temp_examples_dir)
    
    assert len(examples) == 3


def test_load_examples_empty_dir(temp_examples_dir):
    """Тест загрузки из пустой директории"""
    examples = load_examples(temp_examples_dir)
    
    assert len(examples) == 0


def test_load_examples_nonexistent_dir():
    """Тест загрузки из несуществующей директории"""
    examples = load_examples('/nonexistent/directory/12345')
    
    assert len(examples) == 0


def test_load_examples_skips_non_text_files(temp_examples_dir):
    """Тест пропуска файлов не .txt и не .docx"""
    # Создаем файл с другим расширением
    file_path = os.path.join(temp_examples_dir, 'example.pdf')
    with open(file_path, 'w') as f:
        f.write('content')
    
    examples = load_examples(temp_examples_dir)
    
    assert len(examples) == 0


def test_get_relevant_examples_basic():
    """Тест выбора релевантных примеров"""
    description = 'веб-приложение для управления задачами'
    examples = [
        'Веб-приложение для управления проектами и задачами',
        'Мобильное приложение для заметок',
        'Веб-сервис для аналитики данных'
    ]
    
    relevant = get_relevant_examples(description, examples, n_examples=2)
    
    assert len(relevant) == 2
    assert 'Веб-приложение' in relevant[0] or 'Веб-сервис' in relevant[0]


def test_get_relevant_examples_all_if_less():
    """Тест возврата всех примеров, если их меньше запрошенного"""
    description = 'проект'
    examples = ['Пример 1', 'Пример 2']
    
    relevant = get_relevant_examples(description, examples, n_examples=5)
    
    assert len(relevant) == 2


def test_get_relevant_examples_empty_list():
    """Тест выбора из пустого списка"""
    description = 'проект'
    examples = []
    
    relevant = get_relevant_examples(description, examples, n_examples=2)
    
    assert len(relevant) == 0


def test_get_relevant_examples_single_example():
    """Тест выбора одного примера"""
    description = 'веб-приложение'
    examples = ['Веб-приложение для тестирования']
    
    relevant = get_relevant_examples(description, examples, n_examples=1)
    
    assert len(relevant) == 1
    assert relevant[0] == examples[0]


def test_load_examples_handles_unicode_error(temp_examples_dir):
    """Тест обработки ошибок кодировки"""
    # Создаем файл с неверной кодировкой (бинарные данные)
    file_path = os.path.join(temp_examples_dir, 'bad_encoding.txt')
    with open(file_path, 'wb') as f:
        f.write(b'\xff\xfe\x00\x00')  # Невалидный UTF-8
    
    # Должно вернуть пустой список или пропустить файл
    examples = load_examples(temp_examples_dir)
    
    # Файл должен быть пропущен из-за ошибки декодирования
    assert len(examples) == 0
